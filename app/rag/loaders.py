"""Format-specific text extraction. Each loader returns ParsedDocument(s) or raises DocumentLoadError.

Rules: never execute embedded content (HTML scripts are dropped, not run); bounded file size;
malformed documents raise DocumentLoadError (recorded by the caller), never crash the batch.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from app.rag.types import ParsedDocument

MAX_BYTES = 25 * 1024 * 1024
SUPPORTED = {".txt", ".md", ".html", ".htm", ".pdf", ".docx", ".csv", ".xlsx"}


class DocumentLoadError(Exception):
    """A document could not be read or contained no extractable text."""


def _read_bytes(path: Path) -> bytes:
    if path.stat().st_size > MAX_BYTES:
        raise DocumentLoadError(f"{path.name}: exceeds {MAX_BYTES // (1024 * 1024)} MB limit")
    return path.read_bytes()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError("undecodable text encoding")  # pragma: no cover


def load_txt(path: Path) -> list[ParsedDocument]:
    return [ParsedDocument(title=path.stem, text=_decode(_read_bytes(path)))]


def load_html(path: Path) -> list[ParsedDocument]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(_read_bytes(path)), "lxml")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form", "iframe"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else path.stem
    meta: dict[str, object] = {}
    for m in soup.find_all("meta"):
        key = (m.get("property") or m.get("name") or "").lower()
        if key in {"article:published_time", "date", "og:published_time"} and m.get("content"):
            meta["published_raw"] = m["content"]
        if key == "og:url" and m.get("content"):
            meta["canonical_url"] = m["content"]
    body = soup.body or soup
    text = body.get_text(separator="\n")
    return [
        ParsedDocument(
            title=title, text=text, source_url=meta.get("canonical_url"), meta=meta  # type: ignore[arg-type]
        )
    ]


def load_pdf(path: Path) -> list[ParsedDocument]:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(_read_bytes(path)))
        if reader.is_encrypted:
            raise DocumentLoadError(f"{path.name}: encrypted PDF")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except PyPdfError as exc:
        raise DocumentLoadError(f"{path.name}: malformed PDF ({exc})") from exc
    except DocumentLoadError:
        raise
    except Exception as exc:  # noqa: BLE001 - pypdf can raise assorted errors on corrupt input
        raise DocumentLoadError(f"{path.name}: malformed PDF ({type(exc).__name__})") from exc
    text = "\n\n".join(pages).strip()
    if not text:
        raise DocumentLoadError(f"{path.name}: no extractable text (scanned PDF? OCR not enabled)")
    title = reader.metadata.title if reader.metadata and reader.metadata.title else path.stem
    return [ParsedDocument(title=str(title), text=text, meta={"pages": len(pages)})]


def load_docx(path: Path) -> list[ParsedDocument]:
    import docx

    try:
        d = docx.Document(io.BytesIO(_read_bytes(path)))
    except Exception as exc:  # noqa: BLE001
        raise DocumentLoadError(f"{path.name}: malformed DOCX ({type(exc).__name__})") from exc
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [ParsedDocument(title=(d.core_properties.title or path.stem), text="\n".join(parts))]


def load_csv(path: Path) -> list[ParsedDocument]:
    """Each row becomes one small document ('column: value' lines), so rows are individually retrievable."""
    reader = csv.DictReader(io.StringIO(_decode(_read_bytes(path))))
    docs = []
    for i, row in enumerate(reader, start=1):
        text = "\n".join(f"{k}: {v}" for k, v in row.items() if v and str(v).strip())
        if text:
            docs.append(ParsedDocument(title=f"{path.stem} row {i}", text=text, meta={"row": i}))
    return docs


def load_xlsx(path: Path) -> list[ParsedDocument]:
    import pandas as pd

    try:
        sheets = pd.read_excel(io.BytesIO(_read_bytes(path)), sheet_name=None, dtype=str, keep_default_na=False)
    except Exception as exc:  # noqa: BLE001
        raise DocumentLoadError(f"{path.name}: malformed XLSX ({type(exc).__name__})") from exc
    docs = []
    for sheet, df in sheets.items():
        for i, row in enumerate(df.to_dict(orient="records"), start=1):
            text = "\n".join(f"{k}: {v}" for k, v in row.items() if str(v).strip())
            if text:
                docs.append(
                    ParsedDocument(
                        title=f"{path.stem}/{sheet} row {i}", text=text, meta={"sheet": sheet, "row": i}
                    )
                )
    return docs


_LOADERS = {
    ".txt": load_txt, ".md": load_txt, ".html": load_html, ".htm": load_html,
    ".pdf": load_pdf, ".docx": load_docx, ".csv": load_csv, ".xlsx": load_xlsx,
}


def load_document(path: str | Path) -> list[ParsedDocument]:
    p = Path(path)
    if not p.is_file():
        raise DocumentLoadError(f"file not found: {p}")
    loader = _LOADERS.get(p.suffix.lower())
    if loader is None:
        raise DocumentLoadError(f"unsupported type {p.suffix!r}; supported: {sorted(SUPPORTED)}")
    docs = loader(p)
    docs = [d for d in docs if d.text and d.text.strip()]
    if not docs:
        raise DocumentLoadError(f"{p.name}: no text content")
    return docs